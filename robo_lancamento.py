"""
ROBÔ DE LANÇAMENTO DE NOTAS ACADÊMICAS (pyautogui)

Lê o arquivo do cliente (Nome;RA;Nota01;TF01;Nota02;TF02), percorre cada
aluno com um for e digita as notas no Portal do Professor, como um humano.
No fim, confere a planilha gravada pelo sistema contra o arquivo de origem.

Execução:
    1) python sistema_academico.py          (deixe a janela aberta)
    2) python robo_lancamento.py [arquivo.csv]

FAIL-SAFE: arraste o mouse para um CANTO da tela para parar na hora.
"""
import csv
import os
import sys
import time

import pyautogui
from openpyxl import load_workbook

from regras import ler_numero, media_final, situacao

BASE = os.path.dirname(os.path.abspath(__file__))
ARQ_NOTAS = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "dados", "notas_alunos.csv")
ARQ_PLANILHA = os.path.join(BASE, "dados", "notas_lancadas.xlsx")
ARQ_RELATORIO = os.path.join(BASE, "relatorio_execucao.txt")
TITULO_JANELA = "Portal do Professor"

ESPERA_BUSCA = 0.9    # tempo para o sistema "carregar" o aluno (s)
ESPERA_SALVAR = 0.5   # tempo para gravar o registro (s)

pyautogui.FAILSAFE = True   # mouse no canto da tela = parada imediata
pyautogui.PAUSE = 0.05      # pequena pausa entre cada comando


# --------------------------------------------------------- leitura dados ---
def normalizar_coluna(nome):
    # aceita variações do cabeçalho (ex.: "Notao2" do enunciado vira "nota02")
    return nome.strip().lower().replace("notao", "nota0")


def ler_alunos(caminho):
    alunos, erros = [], []
    with open(caminho, encoding="utf-8-sig", newline="") as f:
        leitor = csv.DictReader(f, delimiter=";")
        leitor.fieldnames = [normalizar_coluna(c) for c in leitor.fieldnames]
        for n_linha, linha in enumerate(leitor, start=2):
            try:
                alunos.append({
                    "Nome": linha["nome"].strip(),
                    "RA": linha["ra"].strip(),
                    "Nota01": ler_numero(linha["nota01"]),
                    "TF01": ler_numero(linha["tf01"]),
                    "Nota02": ler_numero(linha["nota02"]),
                    "TF02": ler_numero(linha["tf02"]),
                })
            except (ValueError, KeyError, AttributeError) as erro:
                erros.append(f"Linha {n_linha} ignorada ({erro}): {linha}")
    return alunos, erros


def fmt(valor):
    return f"{valor:g}"  # 7.0 -> "7", 7.5 -> "7.5"


# ------------------------------------------------------------- automação ---
def focar_sistema():
    """Traz a janela do Portal para frente; se não conseguir, pede um clique."""
    try:
        janelas = pyautogui.getWindowsWithTitle(TITULO_JANELA)
        if janelas:
            janela = janelas[0]
            if janela.isMinimized:
                janela.restore()
            janela.activate()
            time.sleep(1)
            return
    except Exception:
        pass
    print("Não achei a janela automaticamente. Clique no campo de BUSCA do Portal.")
    for s in range(5, 0, -1):
        print(f"  Iniciando em {s}...")
        time.sleep(1)


def lancar_aluno(aluno):
    # 1. busca pelo RA
    pyautogui.hotkey("ctrl", "a")
    pyautogui.write(aluno["RA"])
    pyautogui.press("enter")
    time.sleep(ESPERA_BUSCA)

    # 2. notas (o sistema coloca o cursor em "Nota 1" após carregar)
    pyautogui.write(fmt(aluno["Nota01"])); pyautogui.press("tab")
    pyautogui.write(fmt(aluno["TF01"]));   pyautogui.press("tab")
    pyautogui.write(fmt(aluno["Nota02"])); pyautogui.press("tab")
    pyautogui.write(fmt(aluno["TF02"]))

    # 3. salvar (Enter no último campo = botão Salvar)
    pyautogui.press("enter")
    time.sleep(ESPERA_SALVAR)


# ------------------------------------------------------------- conferência ---
def conferir(alunos):
    """Compara a planilha gravada pelo sistema com o arquivo de origem."""
    if not os.path.exists(ARQ_PLANILHA):
        return [f"{a['RA']}: planilha não encontrada" for a in alunos]
    ws = load_workbook(ARQ_PLANILHA)["Notas"]
    cab = [c.value for c in ws[1]]
    gravados = {str(l[0]): dict(zip(cab, l)) for l in ws.iter_rows(min_row=2, values_only=True)}

    divergencias = []
    for a in alunos:
        g = gravados.get(a["RA"])
        if not g:
            divergencias.append(f"{a['RA']} {a['Nome']}: NÃO foi lançado")
            continue
        for campo in ("Nota01", "TF01", "Nota02", "TF02"):
            if abs(float(g[campo]) - a[campo]) > 0.001:
                divergencias.append(f"{a['RA']} {a['Nome']}: {campo} esperado "
                                    f"{a[campo]} / gravado {g[campo]}")
        mf = media_final(a["Nota01"], a["TF01"], a["Nota02"], a["TF02"])
        if g["Situacao"] != situacao(mf):
            divergencias.append(f"{a['RA']} {a['Nome']}: situação esperada "
                                f"{situacao(mf)} / gravada {g['Situacao']}")
    return divergencias


# ------------------------------------------------------------------- main ---
def main():
    alunos, erros_leitura = ler_alunos(ARQ_NOTAS)
    print(f"{len(alunos)} alunos lidos de {os.path.basename(ARQ_NOTAS)}")
    for e in erros_leitura:
        print("  AVISO:", e)
    if not alunos:
        print("Nada para lançar.")
        return

    print("FAIL-SAFE: arraste o mouse para um canto da tela para parar.\n")
    focar_sistema()

    inicio = time.time()
    lancados = 0
    try:
        for i, aluno in enumerate(alunos, start=1):
            lancar_aluno(aluno)
            lancados += 1
            print(f"[{i}/{len(alunos)}] {aluno['RA']} - {aluno['Nome']} lançado")
    except pyautogui.FailSafeException:
        print("\n>>> Robô interrompido pelo FAIL-SAFE (mouse no canto da tela).")
    tempo = time.time() - inicio

    time.sleep(1)
    divergencias = conferir(alunos[:lancados])

    aprovados = sum(1 for a in alunos[:lancados]
                    if situacao(media_final(a["Nota01"], a["TF01"], a["Nota02"], a["TF02"])) == "APROVADO")
    linhas = [
        "RELATÓRIO DE EXECUÇÃO - ROBÔ DE LANÇAMENTO DE NOTAS",
        f"Data: {time.strftime('%d/%m/%Y %H:%M:%S')}",
        f"Arquivo de origem: {os.path.basename(ARQ_NOTAS)}",
        f"Alunos no arquivo: {len(alunos)} | Linhas ignoradas: {len(erros_leitura)}",
        f"Lançados: {lancados} | Aprovados: {aprovados} | Reprovados: {lancados - aprovados}",
        f"Tempo total: {tempo:.1f} s ({tempo / max(lancados, 1):.2f} s por aluno)",
        f"Conferência da planilha: {'OK - zero divergências' if not divergencias else f'{len(divergencias)} divergência(s)'}",
        *[f"  - {d}" for d in divergencias],
        *[f"  - {e}" for e in erros_leitura],
    ]
    with open(ARQ_RELATORIO, "w", encoding="utf-8") as f:
        f.write("\n".join(linhas) + "\n")
    print("\n" + "\n".join(linhas))


if __name__ == "__main__":
    main()
