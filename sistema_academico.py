"""
PORTAL DO PROFESSOR - Sistema Acadêmico Legado (simulado)

Representa o sistema da faculdade que NÃO aceita importação de arquivos:
as notas só entram pela tela, aluno por aluno.

- Busca o aluno por RA ou Nome
- Recebe Nota 1, TF 1, Nota 2, TF 2
- Calcula a média (Nota peso 8, TF peso 2) e a situação
- Salva o registro na planilha dados/notas_lancadas.xlsx

Execução:  python sistema_academico.py
"""
import csv
import os
import tkinter as tk
from datetime import datetime
from tkinter import ttk

from openpyxl import Workbook, load_workbook

from regras import ler_numero, media_bimestre, media_final, situacao

BASE = os.path.dirname(os.path.abspath(__file__))
ARQ_CADASTRO = os.path.join(BASE, "dados", "cadastro_alunos.csv")
ARQ_PLANILHA = os.path.join(BASE, "dados", "notas_lancadas.xlsx")

TITULO_JANELA = "Portal do Professor - Lançamento de Notas"
LATENCIA_BUSCA_MS = 400  # simula o "esperar carregar" do sistema legado

CABECALHO = ["RA", "Nome", "Nota01", "TF01", "Nota02", "TF02",
             "Media01", "Media02", "MediaFinal", "Situacao", "LancadoEm"]


# ----------------------------------------------------------------- dados ---
def carregar_cadastro():
    with open(ARQ_CADASTRO, encoding="utf-8-sig", newline="") as f:
        return [{"RA": l["RA"].strip(), "Nome": l["Nome"].strip()}
                for l in csv.DictReader(f, delimiter=";")]


def abrir_planilha():
    if os.path.exists(ARQ_PLANILHA):
        return load_workbook(ARQ_PLANILHA)
    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"
    ws.append(CABECALHO)
    return wb


def salvar_registro(registro):
    """Grava ou atualiza (pelo RA) a linha do aluno na planilha."""
    wb = abrir_planilha()
    ws = wb["Notas"]
    linha_existente = None
    for linha in range(2, ws.max_row + 1):
        if str(ws.cell(linha, 1).value) == registro["RA"]:
            linha_existente = linha
            break
    valores = [registro[c] for c in CABECALHO]
    if linha_existente:
        for col, v in enumerate(valores, start=1):
            ws.cell(linha_existente, col, v)
    else:
        ws.append(valores)
    wb.save(ARQ_PLANILHA)
    return ws.max_row - 1


# ------------------------------------------------------------------- tela ---
class PortalProfessor(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(TITULO_JANELA)
        self.geometry("560x470")
        self.resizable(False, False)
        self.alunos = carregar_cadastro()
        self.aluno = None
        self.resultados = []
        self._montar_tela()
        self.after(200, self.e_busca.focus_force)

    def _montar_tela(self):
        pad = {"padx": 10, "pady": 4}

        tk.Label(self, text="PORTAL DO PROFESSOR", font=("Segoe UI", 14, "bold"),
                 fg="#1f3a68").pack(pady=(10, 2))
        tk.Label(self, text="Lançamento de notas finais (não permite importação)",
                 fg="#666").pack()

        # --- busca
        fb = tk.LabelFrame(self, text="Buscar aluno (RA ou Nome)")
        fb.pack(fill="x", **pad)
        self.e_busca = ttk.Entry(fb, width=40)
        self.e_busca.grid(row=0, column=0, **pad)
        ttk.Button(fb, text="Buscar", command=self.buscar, takefocus=0).grid(row=0, column=1, **pad)
        self.lst = tk.Listbox(fb, height=3, width=60, takefocus=0)
        self.lst.grid(row=1, column=0, columnspan=2, padx=10, pady=(0, 6))
        self.lst.bind("<<ListboxSelect>>", self._selecionar_lista)

        self.lbl_aluno = tk.Label(self, text="Nenhum aluno selecionado",
                                  font=("Segoe UI", 10, "bold"))
        self.lbl_aluno.pack()

        # --- notas (ordem de criação = ordem do TAB, usada pelo robô)
        fn = tk.LabelFrame(self, text="Notas  (Nota peso 8 | TF peso 2)")
        fn.pack(fill="x", **pad)
        self.campos = []
        for i, rotulo in enumerate(["Nota 1", "TF 1", "Nota 2", "TF 2"]):
            tk.Label(fn, text=rotulo).grid(row=0, column=i, padx=10)
            e = ttk.Entry(fn, width=9, justify="center", state="disabled")
            e.grid(row=1, column=i, padx=10, pady=(0, 8))
            e.bind("<KeyRelease>", lambda _e: self.atualizar_previa())
            self.campos.append(e)
        self.e_n1, self.e_tf1, self.e_n2, self.e_tf2 = self.campos

        self.lbl_previa = tk.Label(self, text="Média final: -   |   Situação: -",
                                   font=("Segoe UI", 10))
        self.lbl_previa.pack(pady=4)

        self.btn_salvar = ttk.Button(self, text="Salvar registro", command=self.salvar,
                                     state="disabled", takefocus=0)
        self.btn_salvar.pack(pady=4)

        self.lbl_status = tk.Label(self, text="Pronto.", fg="#555", anchor="w")
        self.lbl_status.pack(fill="x", side="bottom", padx=10, pady=6)

        # atalhos
        self.e_busca.bind("<Return>", lambda _e: self.buscar())
        self.e_busca.bind("<Control-a>", self._selecionar_tudo)
        for e in self.campos:
            e.bind("<Control-a>", self._selecionar_tudo)
        self.e_tf2.bind("<Return>", lambda _e: self.salvar())

    # --- utilidades
    @staticmethod
    def _selecionar_tudo(evento):
        evento.widget.select_range(0, "end")
        evento.widget.icursor("end")
        return "break"

    def _status(self, texto, cor="#555"):
        self.lbl_status.config(text=texto, fg=cor)

    def _habilitar_campos(self, ativo):
        estado = "normal" if ativo else "disabled"
        for e in self.campos:
            e.config(state=estado)
        self.btn_salvar.config(state=estado)

    def _limpar(self):
        self.aluno = None
        self._habilitar_campos(True)
        for e in self.campos:
            e.delete(0, "end")
        self._habilitar_campos(False)
        self.e_busca.delete(0, "end")
        self.lst.delete(0, "end")
        self.lbl_aluno.config(text="Nenhum aluno selecionado")
        self.lbl_previa.config(text="Média final: -   |   Situação: -", fg="black")
        self.e_busca.focus_set()

    # --- busca
    def buscar(self):
        termo = self.e_busca.get().strip().lower()
        if not termo:
            self._status("Digite um RA ou nome.", "#b00")
            return
        self._status("Carregando...")
        self._habilitar_campos(False)
        self.after(LATENCIA_BUSCA_MS, lambda: self._concluir_busca(termo))

    def _concluir_busca(self, termo):
        exato = [a for a in self.alunos if a["RA"].lower() == termo]
        self.resultados = exato or [a for a in self.alunos
                                    if termo in a["Nome"].lower() or termo in a["RA"].lower()]
        self.lst.delete(0, "end")
        for a in self.resultados:
            self.lst.insert("end", f'{a["RA"]}  -  {a["Nome"]}')
        if not self.resultados:
            self.aluno = None
            self.lbl_aluno.config(text="Nenhum aluno selecionado")
            self._status(f'Aluno "{termo}" não encontrado.', "#b00")
        elif len(self.resultados) == 1:
            self._carregar_aluno(self.resultados[0])
        else:
            self._status(f"{len(self.resultados)} alunos encontrados - selecione na lista.")

    def _selecionar_lista(self, _evento):
        sel = self.lst.curselection()
        if sel:
            self._carregar_aluno(self.resultados[sel[0]])

    def _carregar_aluno(self, aluno):
        self.aluno = aluno
        self.lbl_aluno.config(text=f'{aluno["Nome"]}  |  RA {aluno["RA"]}')
        self._habilitar_campos(True)
        for e in self.campos:
            e.delete(0, "end")
        self.atualizar_previa()
        self._status("Aluno carregado. Informe as notas.")
        self.e_n1.focus_set()

    # --- cálculo e gravação
    def _ler_notas(self):
        return [ler_numero(e.get()) for e in self.campos]

    def atualizar_previa(self):
        try:
            n1, tf1, n2, tf2 = self._ler_notas()
        except ValueError:
            self.lbl_previa.config(text="Média final: -   |   Situação: -", fg="black")
            return
        m = media_final(n1, tf1, n2, tf2)
        sit = situacao(m)
        self.lbl_previa.config(text=f"Média final: {m:.2f}   |   Situação: {sit}",
                               fg="#1a7f37" if sit == "APROVADO" else "#b00")

    def salvar(self):
        if not self.aluno:
            self._status("Nenhum aluno selecionado.", "#b00")
            return
        try:
            n1, tf1, n2, tf2 = self._ler_notas()
        except ValueError as erro:
            self._status(f"Nota inválida: {erro}", "#b00")
            return
        m1, m2 = media_bimestre(n1, tf1), media_bimestre(n2, tf2)
        mf = media_final(n1, tf1, n2, tf2)
        registro = {
            "RA": self.aluno["RA"], "Nome": self.aluno["Nome"],
            "Nota01": n1, "TF01": tf1, "Nota02": n2, "TF02": tf2,
            "Media01": round(m1, 2), "Media02": round(m2, 2),
            "MediaFinal": round(mf, 2), "Situacao": situacao(mf),
            "LancadoEm": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        }
        try:
            total = salvar_registro(registro)
        except PermissionError:
            self._status("Feche a planilha notas_lancadas.xlsx no Excel e salve de novo.", "#b00")
            return
        nome = self.aluno["Nome"]
        self._limpar()
        self._status(f"Salvo: {nome} - {registro['MediaFinal']:.2f} "
                     f"({registro['Situacao']}). Registros na planilha: {total}", "#1a7f37")


if __name__ == "__main__":
    PortalProfessor().mainloop()
